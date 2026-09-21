from pathlib import Path

TEXT_EXTENSIONS = {".md", ".txt", ".csv"}


def extract_text(path: str) -> tuple[str | None, str | None]:
    p = Path(path)
    if not p.exists():
        return None, "파일을 찾을 수 없습니다"
    ext = p.suffix.lower()
    try:
        if ext in {".docx", ".xlsx"} and p.read_bytes()[:8] == bytes.fromhex("d0cf11e0a1b11ae1"):
            return None, "암호화·DRM 또는 구형 OLE 문서입니다. 보호되지 않은 DOCX/XLSX 사본이 필요합니다"
        if ext in TEXT_EXTENSIONS:
            raw = p.read_bytes()
            for enc in ("utf-8", "cp949"):
                try:
                    return raw.decode(enc), None
                except UnicodeDecodeError:
                    continue
            return raw.decode("utf-8", errors="replace"), None
        if ext == ".docx":
            from docx import Document
            doc = Document(p)
            paragraphs = [f"[paragraph {i}] {paragraph.text}" for i, paragraph in enumerate(doc.paragraphs, 1) if paragraph.text]
            rows = [f"[table {i}, row {j}] " + " | ".join(cell.text for cell in row.cells)
                    for i, table in enumerate(doc.tables, 1) for j, row in enumerate(table.rows, 1)]
            return "\n".join(paragraphs + rows), None
        if ext == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(p, read_only=True, data_only=True)
            try:
                return "\n".join(f"[sheet {sheet.title}, row {number}] " + " | ".join(str(v) for v in row if v is not None)
                    for sheet in workbook for number, row in enumerate(sheet.iter_rows(values_only=True), 1) if any(v is not None for v in row)), None
            finally:
                workbook.close()
        if ext == ".pdf":
            from pypdf import PdfReader
            pages = [page.extract_text() or "" for page in PdfReader(p).pages]
            if not any(text.strip() for text in pages):
                return None, "텍스트가 없는 PDF입니다. OCR이 필요한 스캔 문서는 지원하지 않습니다"
            return "\n".join(f"[page {i}]\n{text}" for i, text in enumerate(pages, 1)), None
    except Exception as exc:
        return None, f"문서를 읽지 못했습니다 ({type(exc).__name__})"
    return None, f"{ext or '확장자 없는'} 형식은 아직 내용을 읽을 수 없습니다"
